import io
import pytest
import openpyxl
from flask_jwt_extended import create_access_token
from app import db
from app.models import User, UserRole, Driver, ApprovalStatus, Vehicle, Trip, TripPlan, TripLog, TripStatus
from app.services.routing import ProviderUnavailable


@pytest.fixture
def drivers(app):
    result = []
    for i in range(2):
        user = User(name=f'Driver {i}', email=f'd{i}@example.test', phone=f'55500000{i}', role=UserRole.driver, password_hash='test')
        db.session.add(user); db.session.flush()
        driver = Driver(user_id=user.id, approval_status=ApprovalStatus.approved)
        db.session.add(driver); db.session.flush()
        db.session.add(Vehicle(vehicle_no=f'TEST-{i}', type='Van', fuel_type='Petrol', load_capacity_kg=1000, assigned_driver_id=driver.id))
        db.session.commit()
        result.append({'Authorization': f'Bearer {create_access_token(identity=str(user.id))}'})
    return result


@pytest.fixture
def routes(monkeypatch):
    import app.blueprints.trips.routes as endpoints
    # Route 2 saves emissions and is within the time budget. Route 3 saves
    # more, but is too slow; this tests the actual decision constraint.
    monkeypatch.setattr(endpoints, 'get_road_routes', lambda *_: [
        {'distance_km': 30, 'duration_min': 40, 'pathPoints': [[12.9, 77.5], [12.8, 77.4]]},
        {'distance_km': 25, 'duration_min': 44, 'pathPoints': [[12.9, 77.5], [12.8, 77.4]]},
        {'distance_km': 20, 'duration_min': 60, 'pathPoints': [[12.9, 77.5], [12.8, 77.4]]},
    ])
    monkeypatch.setattr(endpoints, 'estimate', lambda r, *_: {
        'predicted_fuel_l': r['distance_km'] / 10, 'predicted_co2_kg': r['distance_km'],
        'eco_score': 75, 'prediction_mode': 'Test fixture'})


def plan(client, headers):
    res = client.post('/api/v1/trips', headers=headers, json={'source': 'Bengaluru', 'destination': 'Mysuru', 'load_weight_kg': 500})
    assert res.status_code == 201, res.json
    return res.json


def start(client, headers, trip):
    route = next(r for r in trip['routes'] if r['isAiRecommended'])
    assert client.patch(f"/api/v1/trips/{trip['id']}/select-route", headers=headers, json={'route_id': route['id']}).status_code == 200
    res = client.post(f"/api/v1/trips/{trip['id']}/start", headers=headers)
    assert res.status_code == 200, res.json
    return route


def test_complete_loop_survives_reload_and_reconciles(client, drivers, routes, admin_headers):
    trip = plan(client, drivers[0]); tid = trip['id']
    assert trip['status'] == 'planned'
    assert trip['routes'][1]['isAiRecommended']
    assert not trip['routes'][2]['isAiRecommended']
    chosen = start(client, drivers[0], trip)
    loaded = client.get(f'/api/v1/trips/{tid}', headers=drivers[0]).json
    assert loaded['selectedRouteId'] == chosen['id'] and loaded['startedAt']
    assert client.post(f'/api/v1/trips/{tid}/log', headers=drivers[0], json={'lat': 12.85, 'lng': 77.45}).status_code == 201
    fleet = client.get('/api/v1/admin/fleet', headers=admin_headers).json['trips']
    assert fleet[0]['lastLocation']['lat'] == 12.85
    res = client.post(f'/api/v1/trips/{tid}/end', headers=drivers[0], json={'actual_fuel_l': 2.8})
    assert res.status_code == 200 and res.json['completedAt']
    assert res.json['actualCo2Kg'] is None
    assert res.json['estimatedSavingsKg'] == 5
    # Retrying completion cannot overwrite measurements or update totals twice.
    again = client.post(f'/api/v1/trips/{tid}/end', headers=drivers[0], json={'actual_fuel_l': 999})
    assert again.json['actualFuelLiters'] == 2.8
    assert again.json['completedAt'] == res.json['completedAt']
    assert client.get('/api/v1/admin/fleet', headers=admin_headers).json['trips'] == []
    stats = client.get('/api/v1/admin/stats', headers=admin_headers).json
    assert stats['co2_saved_kg'] == 5 and stats['total_trips'] == 1
    assert stats['distance_km'] == 25 and stats['comparable_trips'] == 1
    assert client.get('/api/v1/trips/history', headers=drivers[0]).json['trips'][0]['status'] == 'completed'
    assert client.get('/api/v1/trips/history', headers=drivers[1]).json['trips'] == []
    excel = client.get('/api/v1/reports/export/excel', headers=admin_headers)
    sheet = openpyxl.load_workbook(io.BytesIO(excel.data))['Completed trips']
    assert sheet.cell(2, 7).value == 5
    pdf = client.get('/api/v1/reports/export/pdf', headers=admin_headers)
    assert pdf.status_code == 200 and pdf.data.startswith(b'%PDF')
    assert db.session.get(TripPlan, tid).active_driver_id is None


@pytest.mark.parametrize('method,suffix,body', [
    ('get', '', None), ('patch', '/select-route', {'route_id': 1}),
    ('post', '/start', {}), ('post', '/end', {}), ('post', '/cancel', {}),
    ('post', '/log', {'lat': 12, 'lng': 77}),
])
def test_other_driver_cannot_access_trip(client, drivers, routes, method, suffix, body):
    trip = plan(client, drivers[0])
    res = getattr(client, method)(f"/api/v1/trips/{trip['id']}{suffix}", headers=drivers[1], json=body)
    assert res.status_code == 403


def test_invalid_states_and_reservations(client, drivers, routes):
    trip = plan(client, drivers[0]); other = plan(client, drivers[0]); tid = trip['id']
    assert client.post(f'/api/v1/trips/{tid}/start', headers=drivers[0]).status_code == 409
    assert client.post(f'/api/v1/trips/{tid}/end', headers=drivers[0]).status_code == 409
    assert client.patch(f'/api/v1/trips/{tid}/select-route', headers=drivers[0], json={'route_id': other['routes'][0]['id']}).status_code == 422
    start(client, drivers[0], trip)
    assert client.post(f'/api/v1/trips/{tid}/start', headers=drivers[0]).status_code == 200
    client.patch(f"/api/v1/trips/{other['id']}/select-route", headers=drivers[0], json={'route_id': other['routes'][0]['id']})
    assert client.post(f"/api/v1/trips/{other['id']}/start", headers=drivers[0]).status_code == 409
    assert client.patch(f'/api/v1/trips/{tid}/select-route', headers=drivers[0], json={'route_id': trip['routes'][0]['id']}).status_code == 409
    assert client.post(f'/api/v1/trips/{tid}/cancel', headers=drivers[0]).status_code == 200
    assert client.post(f"/api/v1/trips/{other['id']}/start", headers=drivers[0]).status_code == 200


@pytest.mark.parametrize('weight', [-1, 1001, 'bad', float('inf'), True])
def test_invalid_cargo_is_rejected_without_writes(client, drivers, routes, weight):
    res = client.post('/api/v1/trips', headers=drivers[0], json={'source': 'A', 'destination': 'B', 'load_weight_kg': weight})
    assert res.status_code == 422 and Trip.query.count() == 0


def test_provider_failure_does_not_create_trip(client, drivers, monkeypatch):
    def unavailable(*_): raise ProviderUnavailable('Provider unavailable')
    monkeypatch.setattr('app.blueprints.trips.routes.get_road_routes', unavailable)
    res = client.post('/api/v1/trips', headers=drivers[0], json={'source': 'A', 'destination': 'B'})
    assert res.status_code == 503 and Trip.query.count() == 0


def test_invalid_gps_and_blocked_driver(client, drivers, routes):
    trip = plan(client, drivers[0]); start(client, drivers[0], trip)
    url = f"/api/v1/trips/{trip['id']}/log"
    assert client.post(url, headers=drivers[0], json={'lat': 100, 'lng': 77}).status_code == 422
    assert TripLog.query.count() == 0
    driver = Driver.query.first(); driver.approval_status = ApprovalStatus.blocked; db.session.commit()
    assert client.post(url, headers=drivers[0], json={'lat': 12, 'lng': 77}).status_code == 403


def test_notifications_and_export_permissions(client, drivers, admin_headers):
    from app.models import Notification
    db.session.add(Notification(user_id=Driver.query.first().user_id, type='approval', message='Approved'))
    db.session.commit()
    assert client.get('/api/v1/notifications', headers=drivers[0]).json[0]['title'] == 'Approval'
    assert client.get('/api/v1/notifications', headers=drivers[1]).json == []
    assert client.get('/api/v1/reports/export/excel', headers=drivers[0]).status_code == 403
    assert client.get('/api/v1/reports/export/excel', headers=admin_headers).status_code == 200


def test_json_body_and_unauthenticated_requests(client, drivers):
    assert client.post('/api/v1/trips', json={}).status_code == 401
    assert client.post('/api/v1/trips', headers=drivers[0], json=[]).status_code == 422


def test_unrecognized_location_never_silently_reroutes():
    from app.services.routing import coordinates
    with pytest.raises(ValueError): coordinates('not-a-real-location')
    with pytest.raises(ValueError): coordinates('NaN,77')
    assert coordinates('Bengaluru') == (12.9716, 77.5946)


def test_old_plan_cannot_start(client, drivers, routes):
    from datetime import timedelta
    from app.services.trip_workflow import utcnow
    trip = plan(client, drivers[0])
    client.patch(f"/api/v1/trips/{trip['id']}/select-route", headers=drivers[0], json={'route_id': trip['routes'][0]['id']})
    db.session.get(TripPlan, trip['id']).created_at = utcnow() - timedelta(hours=1)
    db.session.commit()
    assert client.post(f"/api/v1/trips/{trip['id']}/start", headers=drivers[0]).status_code == 409


def test_higher_emissions_are_not_hidden(client, drivers, routes, admin_headers):
    trip = plan(client, drivers[0])
    # Deliberately worse saved alternative to verify signed accounting.
    from app.models import Route
    chosen = trip['routes'][0]
    route = db.session.get(Route, chosen['id']); route.predicted_co2_kg = 40
    stored = db.session.get(TripPlan, trip['id'])
    updated = [dict(r) for r in stored.candidates]; updated[0]['co2Kg'] = 40
    stored.candidates = updated
    db.session.commit()
    client.patch(f"/api/v1/trips/{trip['id']}/select-route", headers=drivers[0], json={'route_id': chosen['id']})
    client.post(f"/api/v1/trips/{trip['id']}/start", headers=drivers[0])
    completed = client.post(f"/api/v1/trips/{trip['id']}/end", headers=drivers[0], json={})
    assert completed.json['estimatedSavingsKg'] == -10
    stats = client.get('/api/v1/admin/stats', headers=admin_headers).json
    assert stats['co2_saved_kg'] == -10


def test_feedback_keeps_zero_reports_and_excludes_demo(client, drivers, routes, admin_headers):
    for reported in (0, 3, None, 100):
        trip = plan(client, drivers[0])
        start(client, drivers[0], trip)
        client.post(f"/api/v1/trips/{trip['id']}/end", headers=drivers[0], json={'actual_fuel_l': reported})
        if reported == 100:
            db.session.get(TripPlan, trip['id']).prediction_mode = 'Demo fixture'
            db.session.commit()
    result = client.get('/api/v1/admin/learning', headers=admin_headers).json
    assert result['eligible_trips'] == 3 and result['reported_trips'] == 2
    assert result['coverage_percent'] == 66.7
    assert result['mae_l'] == 1.5 and result['bias_l'] == -1
    assert len(result['records']) == 2
    assert client.get('/api/v1/admin/learning', headers=drivers[0]).status_code == 403


def test_unsupported_evidence_never_becomes_eco_savings(client, drivers, routes, admin_headers, monkeypatch):
    import app.blueprints.trips.routes as endpoints
    monkeypatch.setattr(endpoints, 'estimate', lambda r, *_: {
        'predicted_fuel_l': r['distance_km'] / 10, 'predicted_co2_kg': r['distance_km'],
        'eco_score': 75, 'prediction_mode': 'Test fixture',
        'evidence': {'within_documented_domain': False}})
    trip = plan(client, drivers[0])
    assert not any(r['isAiRecommended'] for r in trip['routes'])
    default = next(r for r in trip['routes'] if r['isDefault'])
    assert default['isFastest']
    client.patch(f"/api/v1/trips/{trip['id']}/select-route", headers=drivers[0], json={'route_id': default['id']})
    client.post(f"/api/v1/trips/{trip['id']}/start", headers=drivers[0])
    done = client.post(f"/api/v1/trips/{trip['id']}/end", headers=drivers[0], json={}).json
    assert done['estimatedSavingsKg'] is None
    stats = client.get('/api/v1/admin/stats', headers=admin_headers).json
    assert stats['total_trips'] == 1 and stats['comparable_trips'] == 0
    excel = client.get('/api/v1/reports/export/excel', headers=admin_headers)
    sheet = openpyxl.load_workbook(io.BytesIO(excel.data))['Completed trips']
    assert sheet.cell(2, 7).value is None
